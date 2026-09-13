import os
from typing import List, Optional
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, status
from fastapi.responses import FileResponse, PlainTextResponse
from pydantic import BaseModel
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from google import genai

from app.database import get_db
from app.models.document import Document
from app.models.chunk import DocumentChunk
from app.models.user import User
from app.api.auth import get_current_user
from app.config import settings
from app.services.document_service import save_uploaded_file, extract_text_from_file, TextExtractionError

router = APIRouter(prefix="/documents", tags=["Documents"])

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

genai_client = genai.Client(api_key=settings.GEMINI_API_KEY)


class DocumentResponse(BaseModel):
    id: int
    title: str
    file_type: Optional[str] = None
    file_size: Optional[int] = None
    created_at: datetime

    class Config:
        from_attributes = True


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    words = text.split()
    if not words:
        return []
    chunks = []
    for i in range(0, len(words), chunk_size - overlap):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
    return chunks


@router.get("", response_model=List[DocumentResponse])
async def get_user_documents(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    result = await db.execute(select(Document).where(Document.user_id == current_user.id))
    return result.scalars().all()


@router.get("/{document_id}/content")
async def get_document_content(
    document_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    doc = (
        await db.execute(
            select(Document).where(
                Document.id == document_id, Document.user_id == current_user.id
            )
        )
    ).scalar_one_or_none()

    if not doc:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found")

    chunk_rows = (
        await db.execute(
            select(DocumentChunk.content)
            .where(DocumentChunk.document_id == doc.id)
            .order_by(DocumentChunk.chunk_index)
        )
    ).scalars().all()

    full_text = "\n\n".join(chunk_rows) if chunk_rows else "No extracted text available for this document."

    return {"content": full_text}


@router.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    ext = os.path.splitext(file.filename)[1].lower().lstrip(".") if file.filename else "unknown"
    if ext not in ["pdf", "txt", "docx", "doc", "md"]:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file format. Supported formats: .pdf, .docx, .txt, .md",
        )

    # Safe, collision-proof storage name (never trust the original filename for the path)
    stored_filename, file_path, file_size = await save_uploaded_file(file)

    new_doc = Document(
        user_id=current_user.id,
        title=file.filename,
        file_path=file_path,
        file_type=ext,
        file_size=file_size,
    )
    db.add(new_doc)
    await db.commit()
    await db.refresh(new_doc)

    try:
        raw_text = extract_text_from_file(file_path, f".{ext}")
        text_chunks = chunk_text(raw_text)

        for idx, text_content in enumerate(text_chunks):
            if not text_content.strip():
                continue

            emb_res = genai_client.models.embed_content(
                model="gemini-embedding-001",
                contents=text_content,
            )
            vector = emb_res.embeddings[0].values

            chunk_record = DocumentChunk(
                document_id=new_doc.id,
                chunk_index=idx,
                content=text_content,
                embedding=vector,
            )
            db.add(chunk_record)

        await db.commit()
        print(f"[Upload Success] Created {len(text_chunks)} chunks for document ID {new_doc.id}")

    except TextExtractionError as e:
        await db.rollback()
        await db.delete(new_doc)
        await db.commit()
        if os.path.exists(file_path):
            os.remove(file_path)
        print(f"[Upload Rejected] Extraction failed for {file.filename}: {e}")
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=str(e),
        )

    except Exception as e:
        await db.rollback()
        print(f"[Upload Error] Embedding failure: {str(e)}")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to process document embeddings: {str(e)}",
        )

    return {"message": "Document uploaded successfully", "document_id": new_doc.id}


@router.get("/{document_id}/download")
async def download_document(
    document_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    result = await db.execute(
        select(Document).where(Document.id == document_id, Document.user_id == current_user.id)
    )
    doc = result.scalars().first()

    if not doc or not doc.file_path or not os.path.exists(doc.file_path):
        raise HTTPException(status_code=404, detail="Document file not found on server.")

    file_title_lower = doc.title.lower() if doc.title else ""
    if file_title_lower.endswith('.docx'):
        try:
            import docx
            word_doc = docx.Document(doc.file_path)
            full_text = "\n".join([para.text for para in word_doc.paragraphs if para.text.strip()])
            return PlainTextResponse(full_text if full_text else "[Empty Word Document]")
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Failed to parse Word document: {str(e)}")

    return FileResponse(path=doc.file_path, filename=doc.title, media_type="application/octet-stream")


@router.delete("/{document_id}")
async def delete_document(
    document_id: int,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    result = await db.execute(
        select(Document).where(Document.id == document_id, Document.user_id == current_user.id)
    )
    doc = result.scalars().first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")

    if doc.file_path and os.path.exists(doc.file_path):
        try:
            os.remove(doc.file_path)
        except Exception:
            pass

    await db.delete(doc)
    await db.commit()
    return {"message": "Document deleted successfully"}